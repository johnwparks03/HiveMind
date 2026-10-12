"""Text extraction, dispatched by extension. Feature_List 1.3."""
import re
from dataclasses import dataclass
from pathlib import Path

import dateparser
import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation


@dataclass
class ExtractionResult:
    text: str
    in_document_date: "datetime | None"
    failure_reason: str | None = None


def _extract_pdf(path: Path) -> str:
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _extract_docx(path: Path) -> str:
    # Blank line between paragraphs (not just "\n") so chunking.py's
    # paragraph-aware splitter can treat each one as its own chunkable unit.
    doc = DocxDocument(path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(path: Path) -> str:
    # Blank line between slides so each slide can become its own chunk instead
    # of getting pooled with every other slide in the deck (see chunking.py).
    prs = Presentation(path)
    slide_texts = []
    for slide in prs.slides:
        shape_texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if shape_texts:
            slide_texts.append("\n".join(shape_texts))
    return "\n\n".join(slide_texts)


def _extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
}

# A date written in the document body (e.g. a meeting doc's header), distinct from
# filesystem created/modified timestamps. Only the first ~500 chars are checked —
# this is almost always a title/header area.
_DATE_LINE_HINT = re.compile(
    r"(?:date|meeting|notes?)[:\s]*([A-Za-z0-9,./\- ]{6,30})", re.IGNORECASE
)


def _find_in_document_date(text: str):
    head = text[:500]
    match = _DATE_LINE_HINT.search(head)
    candidate = match.group(1).strip() if match else None
    if not candidate:
        return None
    parsed = dateparser.parse(candidate, settings={"STRICT_PARSING": False})
    return parsed


def extract_text(path: Path) -> ExtractionResult:
    ext = path.suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        return ExtractionResult(text="", in_document_date=None, failure_reason=f"unsupported extension {ext}")
    try:
        text = extractor(path)
    except Exception as exc:  # extraction failures are recorded, never crash the agent
        return ExtractionResult(text="", in_document_date=None, failure_reason=str(exc))
    if not text.strip():
        return ExtractionResult(text="", in_document_date=None, failure_reason="no extractable text")
    return ExtractionResult(text=text, in_document_date=_find_in_document_date(text))
