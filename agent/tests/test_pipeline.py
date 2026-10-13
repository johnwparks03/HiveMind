"""Indexing pipeline at corpus scale: every file registered, failures reported, the vector
index kept in step with the registry through rescans, edits and deletions.

These run with fake embeddings (see conftest.py), so they check bookkeeping, not ranking.
"""
from pathlib import Path

import pytest

from app.db import session
from app.folders.scanner import add_folder, remove_folder, scan_folder
from app.indexing.vector_store import get_collection
from corpus_generator import EDGE_CASE_FOLDER, TOTAL_DOCX, TOTAL_PPTX

SUPPORTED_COUNT = TOTAL_DOCX + TOTAL_PPTX + 3  # three edge cases


@pytest.fixture
def indexed_folder(isolated_store, corpus):
    root, _ = corpus
    folder_id = add_folder(str(root))
    return folder_id, root


def _registry_rows(folder_id):
    with session() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM files WHERE folder_id = ?", (folder_id,)).fetchall()]


def _chunks_for(file_id):
    got = get_collection().get(where={"file_id": file_id})
    return got["ids"]


def test_every_supported_file_is_registered(isolated_store, corpus):
    root, _ = corpus
    folder_id = add_folder(str(root))

    rows = _registry_rows(folder_id)
    assert len(rows) == SUPPORTED_COUNT
    assert len({r["path"] for r in rows}) == SUPPORTED_COUNT


def test_only_the_expected_edge_cases_fail(indexed_folder):
    folder_id, _ = indexed_folder
    rows = _registry_rows(folder_id)

    failed = {r["filename"]: r for r in rows if r["index_status"] == "failed"}
    assert set(failed) == {"Blank document.docx", "Corrupt export.docx"}
    for row in failed.values():
        assert row["index_failure_reason"], f"{row['filename']} failed without a reason"

    indexed = [r for r in rows if r["index_status"] == "indexed"]
    assert len(indexed) == SUPPORTED_COUNT - 2
    assert "Heading only.docx" in {r["filename"] for r in indexed}
    assert all(r["index_status"] in ("indexed", "failed") for r in rows)


def test_every_indexed_file_has_chunks_under_its_own_id(indexed_folder):
    folder_id, _ = indexed_folder
    for row in _registry_rows(folder_id):
        if row["index_status"] != "indexed":
            continue
        got = get_collection().get(where={"file_id": row["file_id"]})
        assert got["ids"], f"{row['filename']} is indexed but has no chunks"
        assert {m["filename"] for m in got["metadatas"]} == {row["filename"]}


def test_vector_index_holds_no_orphan_chunks(indexed_folder):
    folder_id, _ = indexed_folder
    file_ids = {r["file_id"] for r in _registry_rows(folder_id)}
    everything = get_collection().get()
    orphans = {m["file_id"] for m in everything["metadatas"]} - file_ids
    assert orphans == set()


def test_rescan_of_unchanged_corpus_touches_nothing(indexed_folder):
    folder_id, root = indexed_folder
    stats = scan_folder(folder_id, str(root))

    assert stats == {"added": 0, "reindexed": 0, "unchanged": SUPPORTED_COUNT, "failed": 0, "removed": 0}


def test_edited_file_is_reindexed_without_duplicating_chunks(indexed_folder):
    folder_id, root = indexed_folder
    target = next(p for p in root.rglob("Budget plan*.docx"))
    row = next(r for r in _registry_rows(folder_id) if r["path"] == str(target))
    before = len(_chunks_for(row["file_id"]))

    from docx import Document
    doc = Document()
    doc.add_heading("Completely different subject", level=1)
    for i in range(12):
        doc.add_paragraph(f"Paragraph {i} about the new subject matter, long enough to force several chunks. " * 4)
    doc.save(target)

    stats = scan_folder(folder_id, str(root))
    assert stats["reindexed"] == 1
    assert stats["unchanged"] == SUPPORTED_COUNT - 1

    got = get_collection().get(where={"file_id": row["file_id"]})
    text = " ".join(got["documents"])
    assert "Paragraph 11" in text, "the new content is not indexed"
    assert "Forecast spend" not in text and "Quarterly budget plan" not in text, "stale text from the old version is still indexed"
    assert before >= 1


def test_deleted_file_leaves_no_registry_row_or_chunks(indexed_folder):
    folder_id, root = indexed_folder
    target = next(root.rglob("Roadmap*.docx"))
    row = next(r for r in _registry_rows(folder_id) if r["path"] == str(target))
    assert _chunks_for(row["file_id"])

    target.unlink()
    stats = scan_folder(folder_id, str(root))

    assert stats["removed"] == 1
    assert all(r["file_id"] != row["file_id"] for r in _registry_rows(folder_id))
    assert _chunks_for(row["file_id"]) == []


def test_file_that_becomes_unreadable_drops_its_old_chunks(indexed_folder):
    """A file that was indexed and then corrupted must not keep serving its old text."""
    folder_id, root = indexed_folder
    target = next(root.rglob("Status report*.docx"))
    row = next(r for r in _registry_rows(folder_id) if r["path"] == str(target))
    assert _chunks_for(row["file_id"])

    target.write_bytes(b"this is no longer a word document")
    stats = scan_folder(folder_id, str(root))

    after = next(r for r in _registry_rows(folder_id) if r["file_id"] == row["file_id"])
    assert after["index_status"] == "failed"
    assert stats["failed"] == 1
    assert _chunks_for(row["file_id"]) == [], "failed file still has searchable chunks from its old content"


def test_removing_a_folder_clears_registry_and_index(indexed_folder):
    folder_id, _ = indexed_folder
    assert get_collection().count() > 0

    assert remove_folder(folder_id) is True

    assert _registry_rows(folder_id) == []
    assert get_collection().count() == 0


def test_chunk_total_matches_sum_of_per_file_chunks(indexed_folder):
    folder_id, _ = indexed_folder
    per_file = sum(len(_chunks_for(r["file_id"])) for r in _registry_rows(folder_id))
    assert per_file == get_collection().count()
    assert per_file > SUPPORTED_COUNT, "expected more than one chunk per file on average"


def test_edge_case_folder_is_scanned(indexed_folder):
    folder_id, root = indexed_folder
    edge_rows = [r for r in _registry_rows(folder_id) if Path(r["path"]).parent.name == EDGE_CASE_FOLDER]
    assert len(edge_rows) == 3
