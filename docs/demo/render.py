"""Renderers for the demo files: DOCX, PPTX, PDF and TXT.

An entry is a dict with: path, title, author, created, revised (YYYY-MM-DD),
and one of blocks (docx, pdf), slides (pptx) or text (txt). See
content/README.md for the block and slide shapes.
"""

import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor
from fpdf import FPDF
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt as PptxPt

# Deloitte brand palette (deloitte-brand skill, colors-typography reference)
GREEN = "86BC25"
NEON = "86EB22"
DARK = "282728"
BLACK = "000000"
WHITE = "FFFFFF"
HEADLINE_FONT = "Open Sans"
ACCENT_FONT = "Stix Two Text"

PDF_REPLACEMENTS = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "•": "-", "…": "...", " ": " ",
}


def ts(date_text):
    return datetime.strptime(date_text, "%Y-%m-%d")


def pdf_safe(text):
    for old, new in PDF_REPLACEMENTS.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def render_docx(entry, out):
    document = Document()
    document.styles["Normal"].font.name = HEADLINE_FONT
    document.styles["Normal"].font.size = Pt(11)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.color.rgb = RGBColor.from_string(DARK)
        document.styles[style_name].font.name = HEADLINE_FONT
    props = document.core_properties
    props.title = entry["title"]
    props.author = entry["author"]
    props.last_modified_by = entry["author"]
    props.created = ts(entry["created"])
    props.modified = ts(entry["revised"])
    for block in entry["blocks"]:
        kind = block[0]
        if kind == "h1":
            document.add_heading(block[1], level=1)
        elif kind == "h2":
            document.add_heading(block[1], level=2)
        elif kind == "p":
            document.add_paragraph(block[1])
        elif kind == "bullets":
            for item in block[1]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            header, rows = block[1], block[2]
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for index, text in enumerate(header):
                cell = table.rows[0].cells[index]
                cell.text = text
                cell.paragraphs[0].runs[0].font.bold = True
            for row in rows:
                cells = table.add_row().cells
                for index, text in enumerate(row):
                    cells[index].text = text
        else:
            raise ValueError(f"Unknown block kind {kind!r} in {entry['path']}")
    document.save(out)


def add_headline(slide, text, dark, top, size):
    box = slide.shapes.add_textbox(Inches(0.6), top, Inches(10.6), Inches(1.4))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    words = text.split()
    body_color = PptxRGB.from_string(WHITE if dark else BLACK)
    if len(words) > 1:
        body_run = paragraph.add_run()
        body_run.text = " ".join(words[:-1]) + " "
        accent_run = paragraph.add_run()
        accent_run.text = words[-1]
        accent_run.font.name = ACCENT_FONT
        accent_run.font.italic = True
        accent_run.font.bold = True
        accent_run.font.color.rgb = PptxRGB.from_string(NEON if dark else BLACK)
        accent_run.font.size = PptxPt(size)
    else:
        body_run = paragraph.add_run()
        body_run.text = text
    body_run.font.name = HEADLINE_FONT
    body_run.font.bold = True
    body_run.font.size = PptxPt(size)
    body_run.font.color.rgb = body_color


def add_bullets(slide, bullets, top):
    box = slide.shapes.add_textbox(Inches(0.6), top, Inches(11.4), Inches(4.8))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = f"• {item}"
        run.font.name = HEADLINE_FONT
        run.font.size = PptxPt(20)
        run.font.color.rgb = PptxRGB.from_string(BLACK)
        paragraph.space_after = PptxPt(12)


def render_pptx(entry, out):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    props = presentation.core_properties
    props.title = entry["title"]
    props.author = entry["author"]
    props.last_modified_by = entry["author"]
    props.created = ts(entry["created"])
    props.modified = ts(entry["revised"])
    for index, content in enumerate(entry["slides"]):
        slide = presentation.slides.add_slide(blank_layout)
        dark = index == 0
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = PptxRGB.from_string(DARK if dark else WHITE)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.7), Inches(0.5), Inches(0.9), Inches(0.9))
        circle.fill.solid()
        circle.fill.fore_color.rgb = PptxRGB.from_string(GREEN)
        circle.line.fill.background()
        if dark:
            add_headline(slide, content["title"], dark=True, top=Inches(2.4), size=44)
            sub = slide.shapes.add_textbox(Inches(0.6), Inches(4.0), Inches(10.6), Inches(1.0))
            sub.text_frame.word_wrap = True
            run = sub.text_frame.paragraphs[0].add_run()
            run.text = content["subtitle"]
            run.font.name = HEADLINE_FONT
            run.font.size = PptxPt(18)
            run.font.color.rgb = PptxRGB.from_string(WHITE)
        else:
            add_headline(slide, content["title"], dark=False, top=Inches(0.6), size=32)
            add_bullets(slide, content["bullets"], top=Inches(2.0))
        footer = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(8), Inches(0.4))
        footer_run = footer.text_frame.paragraphs[0].add_run()
        footer_run.text = "Fictional demo content"
        footer_run.font.name = HEADLINE_FONT
        footer_run.font.size = PptxPt(10)
        footer_run.font.color.rgb = PptxRGB.from_string(WHITE if dark else DARK)
    presentation.save(out)


def render_pdf(entry, out):
    pdf = FPDF()
    pdf.set_creation_date(ts(entry["created"]))
    pdf.set_author(pdf_safe(entry["author"]))
    pdf.set_title(pdf_safe(entry["title"]))
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, 20)
    pdf.add_page()
    for block in entry["blocks"]:
        kind = block[0]
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 8, pdf_safe(block[1]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif kind == "h2":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, pdf_safe(block[1]), new_x="LMARGIN", new_y="NEXT")
        elif kind == "p":
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, pdf_safe(block[1]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "bullets":
            pdf.set_font("Helvetica", "", 11)
            for item in block[1]:
                pdf.multi_cell(0, 6, pdf_safe(f"- {item}"), new_x="LMARGIN", new_y="NEXT")
        elif kind == "table":
            header, rows = block[1], block[2]
            fractions = block[3] if len(block) > 3 else [1 / len(header)] * len(header)
            pdf.set_font("Helvetica", "", 9)
            with pdf.table(col_widths=tuple(fractions), text_align="LEFT") as table:
                head = table.row()
                for text in header:
                    head.cell(pdf_safe(text))
                for row_values in rows:
                    row = table.row()
                    for text in row_values:
                        row.cell(pdf_safe(text))
            pdf.ln(2)
        else:
            raise ValueError(f"Unknown block kind {kind!r} in {entry['path']}")
    pdf.output(out)


def render_txt(entry, out):
    with open(out, "w", encoding="utf-8", newline="\n") as handle:
        handle.write(entry["text"])


RENDERERS = {".docx": render_docx, ".pptx": render_pptx, ".pdf": render_pdf, ".txt": render_txt}


def render_entry(entry, root):
    """Render one entry under root, set its modified time, and return the path."""
    out = Path(root) / entry["path"]
    out.parent.mkdir(parents=True, exist_ok=True)
    RENDERERS[out.suffix.lower()](entry, out)
    revised = ts(entry["revised"]).timestamp()
    os.utime(out, (revised, revised))
    return out
